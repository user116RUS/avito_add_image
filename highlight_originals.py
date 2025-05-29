import pandas as pd
import openpyxl
from openpyxl.styles import PatternFill
import os
import sys

def highlight_original_products(input_file=None, output_file=None):
    """Выделяет все оригинальные товары (без суффикса "-") желтым цветом"""
    
    # Если файлы не указаны, используем значения по умолчанию
    if input_file is None:
        input_file = "few_cities.xlsx"
    
    if output_file is None:
        # Создаем имя выходного файла на основе входного
        name, ext = os.path.splitext(input_file)
        output_file = f"{name}_highlighted{ext}"
    
    print(f"Загрузка Excel-файла: {input_file}")
    
    # Проверяем, существует ли файл
    if not os.path.exists(input_file):
        print(f"Ошибка: Файл '{input_file}' не найден.")
        print("Доступные файлы в текущей директории:")
        files = [f for f in os.listdir('.') if os.path.isfile(f) and (f.endswith('.xlsx') or f.endswith('.xls'))]
        if files:
            for f in files:
                print(f"  - {f}")
        else:
            print("  Нет Excel-файлов (.xlsx или .xls) в текущей директории.")
        return False
    
    try:
        # Загружаем данные из Excel
        df = pd.read_excel(input_file)
        print(f"Загружено {len(df)} строк из Excel")
        
        # Сохраняем DataFrame во временный файл Excel без форматирования
        temp_output = f"temp_{output_file}"
        df.to_excel(temp_output, index=False)
        
        # Открываем Excel-файл с помощью openpyxl для форматирования
        wb = openpyxl.load_workbook(temp_output)
        ws = wb.active
        
        # Создаем заливку желтым цветом
        yellow_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
        
        # Находим индекс столбца с Id
        id_col_index = None
        for i, cell in enumerate(ws[1]):
            if cell.value == 'Id':
                id_col_index = i + 1  # openpyxl использует индексацию с 1
                break
        
        if id_col_index:
            original_count = 0
            derived_count = 0
            
            # Проходим по всем строкам (кроме заголовка)
            for row_idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
                cell = row[id_col_index - 1]  # Получаем ячейку с Id
                
                # Проверяем, является ли это оригинальный товар (без суффикса "-")
                if cell.value and "-" not in str(cell.value):
                    # Заливаем всю строку желтым цветом
                    for cell in row:
                        cell.fill = yellow_fill
                    original_count += 1
                else:
                    derived_count += 1
            
            print(f"Выделено {original_count} оригинальных товаров желтым цветом")
            print(f"Оставлено без выделения {derived_count} производных товаров")
        else:
            print("Ошибка: столбец 'Id' не найден в Excel")
        
        # Сохраняем отформатированный файл
        wb.save(output_file)
        
        # Удаляем временный файл
        os.remove(temp_output)
        
        print(f"Файл успешно сохранен: {output_file}")
        return True
        
    except Exception as e:
        print(f"Ошибка при обработке Excel-файла: {e}")
        import traceback
        traceback.print_exc()
        return False

if __name__ == "__main__":
    # Получаем аргументы командной строки
    input_file = None
    output_file = None
    
    if len(sys.argv) > 1:
        input_file = sys.argv[1]
    
    if len(sys.argv) > 2:
        output_file = sys.argv[2]
    
    # Если аргументы не указаны, ищем любой Excel-файл в текущей директории
    if input_file is None:
        files = [f for f in os.listdir('.') if os.path.isfile(f) and (f.endswith('.xlsx') or f.endswith('.xls'))]
        if files:
            input_file = files[0]
            print(f"Найден Excel-файл: {input_file}")
    
    highlight_original_products(input_file, output_file) 